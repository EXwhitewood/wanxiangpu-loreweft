from __future__ import annotations

import io
import json
import re
import uuid
from datetime import datetime, timezone

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

from app.db.db_models import async_session, Project, Chapter
from sqlalchemy import select
from app.utils.word_count import count_words


async def _load_project_data(project_id: str):
    async with async_session() as session:
        result = await session.execute(
            select(Project).where(Project.id == uuid.UUID(project_id))
        )
        project = result.scalar_one_or_none()
        if not project:
            return None
        chapters_result = await session.execute(
            select(Chapter)
            .where(Chapter.project_id == uuid.UUID(project_id))
            .order_by(Chapter.chapter_number)
        )
        chapters = chapters_result.scalars().all()
        return project, chapters


def _build_markdown(project, chapters):
    core_data = project.core_data or {}
    outline_data = project.outline_data or {}
    story_state = {}
    try:
        from app.services.state_manager import StateManager
        import asyncio
        sm = StateManager()
        state_obj = asyncio.get_event_loop().run_until_complete(sm.get_snapshot(str(project.id)))
        if hasattr(state_obj, "model_dump"):
            story_state = state_obj.model_dump()
    except Exception:
        pass

    lines = []
    lines.append(f"# {project.name}")
    lines.append("")
    if project.description:
        lines.append(f"> {project.description}")
        lines.append("")

    if project.genre:
        lines.append(f"- **题材**: {project.genre}")
    if project.word_count_target:
        lines.append(f"- **目标字数**: {project.word_count_target:,}")
    lines.append(f"- **导出时间**: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    lines.append("")

    outline_chapters = outline_data.get("chapters", [])

    if outline_chapters:
        lines.append("---")
        lines.append("")
        lines.append("## 大纲")
        lines.append("")
        for ch in outline_chapters:
            cn = ch.get("chapter_number", "?")
            title = ch.get("title", "未命名")
            conflict = ch.get("main_conflict", "")
            value_shift = ch.get("value_shift", "")
            pov = ch.get("pov_character", "")
            lines.append(f"### 第{cn}章 {title}")
            if pov:
                lines.append(f"- **视角**: {pov}")
            if conflict:
                lines.append(f"- **核心冲突**: {conflict}")
            if value_shift:
                lines.append(f"- **价值转变**: {value_shift}")
            scenes = ch.get("scenes", [])
            if scenes:
                lines.append("- **场景节拍**:")
                for i, scene in enumerate(scenes):
                    goal = scene.get("goal", "")
                    conflict_s = scene.get("conflict", "")
                    outcome = scene.get("outcome", "")
                    hook = scene.get("hook", "")
                    lines.append(f"  {i+1}. **{goal}**")
                    if conflict_s:
                        lines.append(f"     - 冲突: {conflict_s}")
                    if outcome:
                        lines.append(f"     - 结果: {outcome}")
                    if hook:
                        lines.append(f"     - 钩子: {hook}")
            lines.append("")

    characters = core_data.get("characters", [])

    if characters:
        lines.append("---")
        lines.append("")
        lines.append("## 世界观设定")
        lines.append("")

    if characters:
        lines.append("### 人物卡")
        lines.append("")
        for char in characters[:20]:
            name = char.get("name", "未命名")
            role = char.get("role", "")
            personality = char.get("personality", "")
            desire = char.get("desire", "")
            appearance = char.get("appearance", "")
            lines.append(f"#### {name}")
            if role:
                lines.append(f"- **定位**: {role}")
            if personality:
                lines.append(f"- **性格**: {personality}")
            if desire:
                lines.append(f"- **欲望**: {desire}")
            if appearance:
                lines.append(f"- **外貌**: {appearance}")
            lines.append("")

    if chapters:
        lines.append("---")
        lines.append("")
        lines.append("## 正文")
        lines.append("")
        total_words = 0
        for ch in chapters:
            content = ch.content or ""
            word_count = count_words(content)
            total_words += word_count
            lines.append(f"### 第{ch.chapter_number}章 {ch.title}")
            lines.append("")
            lines.append(f"*字数: {word_count} | 状态: {ch.status}*")
            lines.append("")
            if content.strip():
                lines.append(content)
            else:
                lines.append("*（本章尚未撰写）*")
            lines.append("")
            lines.append("---")
            lines.append("")
        lines.append(f"\n> **全文字数统计: {total_words:,} 字**")

    if story_state:
        active_chapter = story_state.get("active_chapter", "?")
        narrative_time = story_state.get("narrative_time", "—")
        active_characters = story_state.get("active_characters", [])
        active_locations = story_state.get("active_locations", [])

        lines.append("")
        lines.append("---")
        lines.append("")
        lines.append("## 当前状态快照")
        lines.append("")
        lines.append(f"- **当前章节**: 第{active_chapter}章")
        lines.append(f"- **叙事时间**: {narrative_time}")
        if active_characters:
            lines.append(f"- **在场人物**: {', '.join(active_characters)}")
        if active_locations:
            lines.append(f"- **当前地点**: {', '.join(active_locations)}")

    return "\n".join(lines)


def _build_docx(project, chapters):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise RuntimeError("python-docx 未安装")

    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = "SimSun"
    font.size = Pt(12)

    title_style = doc.styles.add_style("DocTitle", 1)
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

    heading2_style = doc.styles["Heading 2"]
    heading2_style.font.color.rgb = RGBColor(0xf5, 0x9e, 0x0b)

    p = doc.add_paragraph(project.name, style="DocTitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if project.description:
        dp = doc.add_paragraph(project.description)
        dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        dp.runs[0].font.italic = True
        dp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if project.genre:
        info_p.add_run(f"题材: {project.genre}  ")
    info_p.add_run(f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    doc.add_page_break()

    core_data = project.core_data or {}
    outline_data = project.outline_data or {}
    outline_chapters = outline_data.get("chapters", [])
    characters = core_data.get("characters", [])

    if outline_chapters:
        h2 = doc.add_heading("大纲", level=2)
        for ch in outline_chapters:
            cn = ch.get("chapter_number", "?")
            title = ch.get("title", "未命名")
            h3 = doc.add_heading(f"第{cn}章 {title}", level=3)
            conflict = ch.get("main_conflict", "")
            value_shift = ch.get("value_shift", "")
            if conflict:
                doc.add_paragraph(f"核心冲突: {conflict}")
            if value_shift:
                doc.add_paragraph(f"价值转变: {value_shift}")
            scenes = ch.get("scenes", [])
            for si, scene in enumerate(scenes):
                goal = scene.get("goal", "")
                outcome = scene.get("outcome", "")
                text = f"场景{si + 1}: {goal}"
                if outcome:
                    text += f" → {outcome}"
                doc.add_paragraph(text, style="List Bullet")

    if characters:
        doc.add_heading("人物卡", level=2)
        for char in characters[:15]:
            name = char.get("name", "未命名")
            role = char.get("role", "")
            personality = char.get("personality", "")
            p = doc.add_paragraph()
            run = p.add_run(name)
            run.bold = True
            run.font.size = Pt(13)
            if role:
                p.add_run(f" ({role})")
            if personality:
                doc.add_paragraph(personality, style="List Bullet")

    doc.add_page_break()

    doc.add_heading("正文", level=2)
    total_words = 0
    for ch in chapters:
        content = ch.content or ""
        wc = count_words(content)
        total_words += wc
        ch_h = doc.add_heading(f"第{ch.chapter_number}章 {ch.title}", level=3)
        meta = ch_h.add_run(f"  ({wc}字)")
        meta.font.size = Pt(10)
        meta.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        if content.strip():
            for para_text in content.split("\n"):
                if para_text.strip():
                    doc.add_paragraph(para_text)
                else:
                    doc.add_paragraph("")
        else:
            empty_p = doc.add_paragraph("（本章尚未撰写）")
            empty_p.runs[0].font.italic = True
            empty_p.runs[0].font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    summary_p = doc.add_paragraph()
    summary_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = summary_p.add_run(f"全文字数: {total_words:,} 字")
    run.bold = True
    run.font.size = Pt(11)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def parse_markdown_import(md_content: str) -> dict:
    result = {
        "project_name": "",
        "description": "",
        "genre": "",
        "chapters": [],
        "outline_chapters": [],
        "characters": [],
    }

    lines = md_content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        title_match = re.match(r"^#\s+(.+)$", line)
        if title_match and not result["project_name"]:
            result["project_name"] = title_match.group(1).strip()
            i += 1
            continue

        blockquote_match = re.match(r"^>\s*(.+)$", line)
        if blockquote_match and not result["description"]:
            desc_lines = [blockquote_match.group(1).strip()]
            j = i + 1
            while j < len(lines) and lines[j].lstrip().startswith(">"):
                desc_lines.append(lines[j].lstrip()[1:].strip())
                j += 1
            result["description"] = "\n".join(desc_lines)
            i = j
            continue

        genre_match = re.match(r"^-\s*\*\*题材\*\*:\s*(.+)$", line)
        if genre_match:
            result["genre"] = genre_match.group(1).strip()
            i += 1
            continue

        chapter_header = re.match(r"^###\s*第(\d+)章\s*(.*)$", line)
        if chapter_header:
            ch_num = int(chapter_header.group(1))
            ch_title = chapter_header.group(2).strip() or f"第{ch_num}章"
            content_lines = []
            i += 1
            while i < len(lines):
                next_line = lines[i].rstrip()
                if re.match(r"^#{1,3}\s", next_line) or next_line == "---":
                    break
                if next_line.startswith("- **") or next_line.startswith("*") or next_line.startswith(">"):
                    break
                content_lines.append(next_line)
                i += 1
            content = "\n".join(content_lines).strip()
            if content and len(content) > 10:
                result["chapters"].append({
                    "chapter_number": ch_num,
                    "title": ch_title,
                    "content": content,
                    "status": "draft",
                })
            else:
                result["outline_chapters"].append({
                    "chapter_number": ch_num,
                    "title": ch_title,
                })
            continue

        char_header = re.match(r"^####\s+(.+)$", line)
        if char_header:
            char_name = char_header.group(1).strip()
            char_data = {"name": char_name}
            i += 1
            while i < len(lines):
                next_line = lines[i].rstrip()
                field_match = re.match(r"^-\s*\*\*(\w+)\*\*:\s*(.+)$", next_line)
                if field_match:
                    key = field_match.group(1).lower()
                    val = field_match.group(2).strip()
                    if key in ("定位", "role"):
                        char_data["role"] = val
                    elif key in ("性格", "personality"):
                        char_data["personality"] = val
                    elif key in ("欲望", "desire"):
                        char_data["desire"] = val
                    elif key in ("外貌", "appearance"):
                        char_data["appearance"] = val
                    i += 1
                else:
                    break
            result["characters"].append(char_data)
            continue
        i += 1

    if not result["project_name"] and result["chapters"]:
        result["project_name"] = "导入项目"
    return result


def parse_docx_import(file_bytes: bytes) -> dict:
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx 未安装")

    buf = io.BytesIO(file_bytes)
    doc = Document(buf)

    paragraphs = doc.paragraphs
    result = {
        "project_name": "",
        "chapters": [],
        "outline_chapters": [],
    }

    for i, p in enumerate(paragraphs):
        text = p.text.strip()
        if not text:
            continue
        if i == 0 and len(text) >= 2 and len(text) <= 50 and not any(c in text for c in [".", ",", "。", "，"]):
            result["project_name"] = text
            break

    current_chapter = None
    current_content_lines = []
    chapter_pattern = re.compile(r"^第(\d+)章\s*(.*)$")

    for p in paragraphs:
        text = p.text
        m = chapter_pattern.match(text.strip())
        if m:
            if current_chapter and current_content_lines:
                content = "\n".join(current_content_lines).strip()
                if len(content) > 20:
                    result["chapters"].append({
                        "chapter_number": current_chapter["num"],
                        "title": current_chapter["title"],
                        "content": content,
                        "status": "draft",
                    })
                else:
                    result["outline_chapters"].append({
                        "chapter_number": current_chapter["num"],
                        "title": current_chapter["title"],
                    })
            current_chapter = {"num": int(m.group(1)), "title": m.group(2).strip() or f"第{m.group(1)}章"}
            current_content_lines = []
        elif current_chapter:
            clean_text = text.rstrip()
            if clean_text and not re.match(r"^\d+\s*字$", clean_text):
                current_content_lines.append(clean_text)

    if current_chapter and current_content_lines:
        content = "\n".join(current_content_lines).strip()
        if len(content) > 20:
            result["chapters"].append({
                "chapter_number": current_chapter["num"],
                "title": current_chapter["title"],
                "content": content,
                "status": "draft",
            })
        else:
            result["outline_chapters"].append({
                "chapter_number": current_chapter["num"],
                "title": current_chapter["title"],
            })

    if not result["project_name"] and result["chapters"]:
        result["project_name"] = "导入项目"
    return result


router = APIRouter()


class ProjectArchivePreviewRequest(BaseModel):
    archive: dict


class ProjectArchiveRestoreRequest(BaseModel):
    archive: dict
    confirmation: str


@router.get("/archive", summary="导出无损项目归档")
async def export_project_archive(project_id: str):
    from app.services.project_archive_service import ProjectArchiveService

    async with async_session() as session:
        try:
            archive = await ProjectArchiveService().export_project(session, project_id)
        except ValueError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
    content = json.dumps(archive, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return Response(
        content=content,
        media_type="application/vnd.loreweft.project+json",
        headers={
            "Content-Disposition": f'attachment; filename="{project_id}.loreweft.json"',
            "X-Archive-Checksum": str((archive.get("checksum") or {}).get("tables") or ""),
        },
    )


@router.post("/archive/preview", summary="预检无损项目归档")
async def preview_project_archive(project_id: str, data: ProjectArchivePreviewRequest):
    from app.services.project_archive_service import ProjectArchiveService

    async with async_session() as session:
        try:
            return await ProjectArchiveService().preview_restore(
                session,
                project_id=project_id,
                archive=data.archive,
            )
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post("/archive/restore", summary="事务恢复无损项目归档")
async def restore_project_archive(project_id: str, data: ProjectArchiveRestoreRequest):
    from app.services.project_archive_service import ProjectArchiveService

    async with async_session() as session:
        try:
            return await ProjectArchiveService().restore_project(
                session,
                project_id=project_id,
                archive=data.archive,
                confirmation=data.confirmation,
            )
        except ValueError as exc:
            await session.rollback()
            raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/markdown", summary="导出为 Markdown")
async def export_markdown(project_id: uuid.UUID):
    import asyncio

    data = await _load_project_data(str(project_id))
    if not data:
        raise HTTPException(status_code=404, detail="项目不存在")

    md_content = await asyncio.to_thread(_build_markdown, data[0], data[1])
    return Response(
        content=md_content.encode("utf-8"),
        media_type="text/markdown; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="loreweft-{str(project_id)[:8]}.md"'
        },
    )


@router.post("/import-markdown", summary="从 Markdown 导入")
async def import_markdown(project_id: uuid.UUID, data: dict):
    import asyncio

    content = data.get("content", "")
    if not content.strip():
        raise HTTPException(status_code=400, detail="Markdown 内容不能为空")

    parsed = await asyncio.to_thread(parse_markdown_import, content)

    async with async_session() as db:
        project = await db.get(Project, project_id)
        if not project:
            raise HTTPException(status_code=404, detail="项目不存在")

        imported = {"chapters": 0, "characters": 0, "outline_updated": False, "conflicts_skipped": 0}

        if parsed.get("project_name") and not project.name:
            project.name = parsed["project_name"]
        if parsed.get("description") and not project.description:
            project.description = parsed["description"]
        if parsed.get("genre") and not project.genre:
            project.genre = parsed["genre"]

        characters = parsed.get("characters", [])
        if characters:
            core_data = project.core_data or {}
            existing_chars = core_data.get("characters", [])
            existing_names = {c.get("name", "") for c in existing_chars}
            now = datetime.now(timezone.utc).isoformat()
            for char in characters:
                if char.get("name") and char["name"] not in existing_names:
                    char["id"] = str(uuid.uuid4())
                    char["project_id"] = str(project_id)
                    char["created_at"] = now
                    char["updated_at"] = now
                    existing_chars.append(char)
                    imported["characters"] += 1
            core_data["characters"] = existing_chars
            project.core_data = core_data

        outline_chapters = parsed.get("outline_chapters", [])
        if outline_chapters and not project.outline_data:
            from app.services.story_plan_service import StoryPlanService
            sp_service = StoryPlanService()
            outline_data = {
                "chapters": [{"chapter_number": c["chapter_number"], "title": c["title"]} for c in outline_chapters],
            }
            await sp_service.save_outline_data(project_id, outline_data, db, mode="replace")
            imported["outline_updated"] = True

        for ch_data in parsed.get("chapters", []):
            ch_num = ch_data["chapter_number"]
            result = await db.execute(
                select(Chapter).where(
                    Chapter.project_id == project_id,
                    Chapter.chapter_number == ch_num,
                )
            )
            existing = result.scalar_one_or_none()
            if existing:
                if ch_data.get("content") and ch_data["content"] != (existing.content or ""):
                    imported["conflicts_skipped"] += 1
            else:
                db.add(Chapter(
                    project_id=project_id,
                    chapter_number=ch_num,
                    title=ch_data.get("title", f"第{ch_num}章"),
                    content=ch_data.get("content", ""),
                    status=ch_data.get("status", "draft"),
                ))
                imported["chapters"] += 1

        from app.services.project_chapter_aggregate_service import refresh_project_chapter_aggregates
        await refresh_project_chapter_aggregates(db, project_id=project_id, project=project)

        await db.commit()

    return {
        "message": "Markdown 导入完成",
        "imported": imported,
        "parsed_info": {
            "project_name": parsed.get("project_name", ""),
            "chapters_found": len(parsed.get("chapters", [])),
            "outline_chapters_found": len(parsed.get("outline_chapters", [])),
            "characters_found": len(parsed.get("characters", [])),
        },
    }


@router.get("/docx", summary="导出为 Word (.docx)")
async def export_docx(project_id: uuid.UUID):
    import asyncio

    data = await _load_project_data(str(project_id))
    if not data:
        raise HTTPException(status_code=404, detail="项目不存在")

    buf = await asyncio.to_thread(_build_docx, data[0], data[1])
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="loreweft-{str(project_id)[:8]}.docx"'
        },
    )


@router.post("/import-docx", summary="从 Word 导入")
async def import_docx(project_id: uuid.UUID, data: dict):
    import asyncio
    import base64

    file_b64 = data.get("file_base64", "")
    if not file_b64:
        raise HTTPException(status_code=400, detail="文件数据不能为空")

    try:
        file_bytes = base64.b64decode(file_b64)
    except Exception:
        raise HTTPException(status_code=400, detail="文件数据格式错误")

    parsed = await asyncio.to_thread(parse_docx_import, file_bytes)

    async with async_session() as db:
        project = await db.get(Project, project_id)
        if not project:
            raise HTTPException(status_code=404, detail="项目不存在")

        imported = {"chapters": 0, "outline_updated": False, "conflicts_skipped": 0}
        if parsed.get("project_name") and not project.name:
            project.name = parsed["project_name"]

        for ch_data in parsed.get("chapters", []):
            ch_num = ch_data["chapter_number"]
            result = await db.execute(
                select(Chapter).where(
                    Chapter.project_id == project_id,
                    Chapter.chapter_number == ch_num,
                )
            )
            existing = result.scalar_one_or_none()
            if existing:
                if ch_data.get("content") and ch_data["content"] != (existing.content or ""):
                    imported["conflicts_skipped"] += 1
            else:
                db.add(Chapter(
                    project_id=project_id,
                    chapter_number=ch_num,
                    title=ch_data.get("title", f"第{ch_num}章"),
                    content=ch_data.get("content", ""),
                    status="draft",
                ))
                imported["chapters"] += 1

        outline_chapters = parsed.get("outline_chapters", [])
        if outline_chapters and not project.outline_data:
            from app.services.story_plan_service import StoryPlanService
            sp_service = StoryPlanService()
            outline_data = {
                "chapters": [{"chapter_number": c["chapter_number"], "title": c["title"]} for c in outline_chapters],
            }
            await sp_service.save_outline_data(project_id, outline_data, db, mode="replace")
            imported["outline_updated"] = True

        from app.services.project_chapter_aggregate_service import refresh_project_chapter_aggregates
        await refresh_project_chapter_aggregates(db, project_id=project_id, project=project)

        await db.commit()

    return {
        "message": "Word 文档导入完成",
        "imported": imported,
        "parsed_info": {
            "project_name": parsed.get("project_name", ""),
            "chapters_found": len(parsed.get("chapters", [])),
            "outline_chapters_found": len(parsed.get("outline_chapters", [])),
        },
    }
